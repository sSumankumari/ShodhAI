from flask import (
    Flask, render_template, request, jsonify, send_file, redirect, url_for, make_response
)
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
import os
import io
import tempfile

# Load environment variables
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

# Import feature modules
from features.ocr_extractor import extract_text_from_file
from features.plagiarism_checker import check_plagiarism_from_files, check_plagiarism_from_strings
from features.rag_pdf_chatbot import answer_from_pdf
from features.webURL_analyzer import analyze as web_analyze
from features.youtube_analyzer import (
    extract_and_save_transcript, ask_question_over_transcript, summarize_transcript
)

# Groq/OpenAI client (if needed)
try:
    from groq import Groq
    groq_client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None
except ImportError:
    groq_client = None

app = Flask(__name__, template_folder='templates', static_folder='static')

# ========== API ROUTES FOR FRONTEND AJAX ==========

@app.route('/api/ocr', methods=['POST'])
def api_ocr():
    file = request.files.get('image')
    if not file or file.filename == '':
        return jsonify({'success': False, 'error': "No file selected."}), 400
    try:
        file.stream.seek(0)
        text = extract_text_from_file(file.stream)
        return jsonify({'success': True, 'extracted_text': text})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/ocr/download', methods=['POST'])
def api_ocr_download():
    data = request.json
    text = data.get('text', '')
    fmt = data.get('format', 'txt')
    filename = f"extracted_text.{fmt}"
    try:
        if fmt == 'txt':
            output = io.BytesIO(text.encode('utf-8'))
            return send_file(output, as_attachment=True, download_name=filename, mimetype='text/plain')
        elif fmt == 'pdf':
            from fpdf import FPDF
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font("Arial", size=12)
            for line in text.split("\n"):
                pdf.cell(0, 10, line, ln=True)
            out = io.BytesIO()
            pdf.output(out)
            out.seek(0)
            return send_file(out, as_attachment=True, download_name=filename, mimetype='application/pdf')
        elif fmt == 'docx':
            from docx import Document
            doc = Document()
            doc.add_paragraph(text)
            out = io.BytesIO()
            doc.save(out)
            out.seek(0)
            return send_file(out, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        else:
            return jsonify({'success': False, 'error': "Invalid format"}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/plagiarism', methods=['POST'])
def api_plagiarism():
    files = request.files.getlist('files[]')
    if len(files) < 2:
        return jsonify({'success': False, 'error': "Please upload at least 2 files."}), 400
    file_names = [f.filename for f in files]
    try:
        results = check_plagiarism_from_files(files, file_names)
        return jsonify({'success': True, 'results': results, 'file_names': file_names})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/plagiarism/download', methods=['POST'])
def api_plagiarism_download():
    data = request.json
    table = data.get('table')
    filename = data.get('filename', 'results.csv')
    fmt = data.get('format', 'csv')
    try:
        if fmt == 'csv':
            output = io.StringIO()
            # The frontend should send CSV content as `table`
            output.write(table)
            output.seek(0)
            return send_file(io.BytesIO(output.getvalue().encode('utf-8')), as_attachment=True, download_name=filename, mimetype='text/csv')
        elif fmt == 'html':
            # The frontend should send HTML table string as `table`
            output = io.StringIO()
            output.write(table)
            output.seek(0)
            return send_file(io.BytesIO(output.getvalue().encode('utf-8')), as_attachment=True, download_name=filename, mimetype='text/html')
        else:
            return jsonify({'success': False, 'error': "Invalid format"}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/rag_pdf/upload', methods=['POST'])
def api_rag_pdf_upload():
    pdf = request.files.get('pdf')
    if not pdf:
        return jsonify({'success': False, 'error': "PDF required."}), 400
    # Store temporarily (simulate session, or return a temp token, or handle in memory)
    try:
        temp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        pdf.save(temp.name)
        temp.close()
        return jsonify({'success': True, 'temp_pdf': temp.name})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/rag_pdf/ask', methods=['POST'])
def api_rag_pdf_ask():
    temp_pdf = request.form.get('temp_pdf')
    question = request.form.get('question', '')
    if not temp_pdf or not question:
        return jsonify({'success': False, 'error': "PDF and question required."}), 400
    try:
        with open(temp_pdf, 'rb') as f:
            answer = answer_from_pdf(f, question, temperature=0.2)
        return jsonify({'success': True, 'answer': answer})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/web_analyze', methods=['POST'])
def api_web_analyze():
    data = request.json
    url = data.get('url')
    question = data.get('question', '')
    if not url:
        return jsonify({'success': False, 'error': "URL required."}), 400
    try:
        result = web_analyze(url, question, GROQ_API_KEY)
        if "error" in result:
            return jsonify({'success': False, 'error': result["error"]}), 400
        return jsonify({'success': True, 'result': result})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/youtube/extract', methods=['POST'])
def api_youtube_extract():
    data = request.json
    youtube_url = data.get('url')
    if not youtube_url:
        return jsonify({'success': False, 'error': "YouTube URL required."}), 400
    result = extract_and_save_transcript(youtube_url)
    if result['success']:
        return jsonify({'success': True, 'transcript': result['transcript']})
    else:
        return jsonify({'success': False, 'error': result['error']}), 400

@app.route('/api/youtube/ask', methods=['POST'])
def api_youtube_ask():
    data = request.json
    question = data.get('question', '')
    if not question:
        return jsonify({'success': False, 'error': "Question required."}), 400
    result = ask_question_over_transcript(question, client=groq_client)
    if result['success']:
        return jsonify({'success': True, 'answer': result['answer']})
    else:
        return jsonify({'success': False, 'error': result['error']}), 400

@app.route('/api/youtube/summarize', methods=['POST'])
def api_youtube_summarize():
    data = request.json
    transcript = data.get('transcript', '')
    if not transcript:
        return jsonify({'success': False, 'error': "Transcript required."}), 400
    result = summarize_transcript(transcript, client=groq_client)
    if result['success']:
        return jsonify({'success': True, 'summary': result['summary']})
    else:
        return jsonify({'success': False, 'error': result['error']}), 400

# ========== STATIC FILES AND MAIN ROUTES ==========

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/static/<path:filename>')
def static_files(filename):
    return send_file(os.path.join(app.static_folder, filename))

# ========== ERROR HANDLERS ==========

@app.errorhandler(413)
def too_large(e):
    return "File is too large", 413

@app.errorhandler(404)
def not_found(e):
    return render_template('404.html'), 404

@app.errorhandler(500)
def internal_error(e):
    return render_template('500.html'), 500

# ========== MAIN ==========

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(debug=True, host='0.0.0.0', port=port)